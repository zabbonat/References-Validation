"""
Build the two manual coding sheets for the reference audit.

Sheet 1: census of all unresolved (flagged) strings.
Sheet 2: proportional stratified random sample of the automatically verified
         stratum, used to bound the rate at which a fabricated reference is
         labelled Verified by the automated pass.

Both sheets are written empty. Every verdict column is filled by a human coder.
The sample is reproducible from SEED.
"""

import urllib.parse

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20260726
N_VERIFIED_SAMPLE = 200
DOUBLE_CODE_SHARE = 0.25

DATA = r"C:\Users\Dilet\Desktop\References-Validation\reviewers-copilot\data\longitudinal_dataset_N22479.xlsx"
OUT_DIR = r"C:\Users\Dilet\Desktop\Scientometrics_Results"

COHORT = {
    "Arxiv_Bio_2016": "2016", "Arxiv_CS_2016": "2016", "NeurIPS_2016": "2016",
    "Arxiv_Bio_2023": "2023", "Arxiv_CS_2023": "2023", "NeurIPS_2023": "2023",
    "Arxiv_CS_2025": "2025", "NeurIPS_2025": "2025",
    "Arxiv_General_Latest": "2026",
}

FLAG_CATEGORIES = [
    "Extraction artifact",
    "Coverage limitation",
    "Parser error",
    "Fabrication",
    "Undecided",
]
EXISTS_CATEGORIES = ["Real work", "Fabrication", "Undecided"]
MATCH_CATEGORIES = ["Correct match", "Wrong record matched", "Undecided"]

df = pd.read_excel(DATA)
df["Cohort"] = df["Dataset"].map(COHORT)

flagged = df[df["Status"] == "Not Found"].copy()
verified = df[df["Status"] == "Verified"].copy()
print(f"flagged: {len(flagged)}   verified: {len(verified)}")

# proportional allocation across the nine venue-year strata, largest remainder
shares = verified["Dataset"].value_counts(normalize=True)
exact = shares * N_VERIFIED_SAMPLE
alloc = exact.apply(int)
while alloc.sum() < N_VERIFIED_SAMPLE:
    alloc[(exact - alloc).idxmax()] += 1
    exact[(exact - alloc).idxmax()] -= 1e-9

sample = pd.concat(
    verified[verified["Dataset"] == d].sample(n=int(k), random_state=SEED)
    for d, k in alloc.items()
)
# shuffle so that any prefix of the sheet remains an unbiased subsample
sample = sample.sample(frac=1.0, random_state=SEED).reset_index(drop=True)
flagged = flagged.sample(frac=1.0, random_state=SEED).reset_index(drop=True)

print(alloc.to_string())


def scholar_url(ref):
    q = urllib.parse.quote_plus(str(ref)[:120])
    return f"https://scholar.google.com/scholar?q={q}"


HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")


def write_sheet(ws, rows, headers, input_cols, widths):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for r in rows:
        ws.append(r)

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ref_col = headers.index("Original Reference") + 1
    for row in range(2, len(rows) + 2):
        ws.cell(row=row, column=ref_col).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 42
        for c in input_cols:
            ws.cell(row=row, column=headers.index(c) + 1).fill = INPUT_FILL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def add_dropdown(ws, headers, col_name, options, n_rows):
    dv = DataValidation(
        type="list", formula1='"' + ",".join(options) + '"', allow_blank=True, showDropDown=False
    )
    dv.error = "Select a value from the list."
    dv.errorTitle = "Invalid entry"
    ws.add_data_validation(dv)
    letter = get_column_letter(headers.index(col_name) + 1)
    dv.add(f"{letter}2:{letter}{n_rows + 1}")


# ----------------------------------------------------------------- sheet 1
wb = Workbook()
ws = wb.active
ws.title = "Coding"

headers1 = [
    "ID", "Cohort", "Dataset", "Paper File", "Original Reference", "Search",
    "Verdict", "Evidence (URL or DOI)", "Indexed in", "Notes", "Double_Coded", "Coder",
]
rows1 = []
n_dbl = int(round(len(flagged) * DOUBLE_CODE_SHARE))
dbl_idx = set(flagged.sample(n=n_dbl, random_state=SEED).index)
for i, (idx, r) in enumerate(flagged.iterrows(), start=1):
    ref = str(r["Original Reference"])
    rows1.append([
        f"F{i:03d}", r["Cohort"], r["Dataset"], r["Paper File"], ref,
        f'=HYPERLINK("{scholar_url(ref)}","search")',
        None, None, None, None,
        "YES" if idx in dbl_idx else "NO", None,
    ])

write_sheet(ws, rows1, headers1,
            input_cols=["Verdict", "Evidence (URL or DOI)", "Indexed in", "Notes", "Coder"],
            widths=[8, 8, 20, 30, 70, 9, 22, 30, 22, 40, 13, 10])
add_dropdown(ws, headers1, "Verdict", FLAG_CATEGORIES, len(rows1))
add_dropdown(ws, headers1, "Indexed in",
             ["Crossref", "OpenAlex", "Semantic Scholar", "arXiv/DBLP", "None of the four"],
             len(rows1))

readme = wb.create_sheet("Protocol")
for line in [
    ["AUDIT OF UNRESOLVED STRINGS - CODING SHEET"],
    [""],
    [f"Rows: {len(rows1)} (census of every string the automated pass left unresolved)"],
    ["Row order randomised; the sheet carries no information in its ordering."],
    [""],
    ["PROCEDURE, per row"],
    ["1. Read the string. Decide first whether it is a bibliographic reference at all."],
    ["   If it is a formula, a heading, a funding note or other layout text -> Parser error."],
    ["2. If it is a reference, search for the cited work. Use the Search link, then widen"],
    ["   the search by hand if it fails: title alone, author + year, exact-phrase, the"],
    ["   publisher site, the author's own page."],
    ["3. If the work is found, record where in 'Evidence'. Then decide between:"],
    ["   Extraction artifact  - the work is indexed in at least one of the four structured"],
    ["                          sources, and the automated pass failed only because the"],
    ["                          extracted string was corrupted."],
    ["   Coverage limitation  - the work exists but none of the four sources indexes it."],
    ["   Use the CheckIfExist dashboard to settle this, and record the answer in 'Indexed in'."],
    ["4. If the work cannot be found after a genuine search -> Fabrication."],
    ["5. If you cannot decide, use Undecided and say why in Notes. Undecided is a valid"],
    ["   outcome and is more useful than a forced call."],
    [""],
    ["DOUBLE CODING"],
    ["Rows marked YES in Double_Coded are coded independently by a second coder, who must"],
    ["not see the first coder's column. Copy the sheet, delete the filled columns, and give"],
    ["the copy to the second coder. Cohen's kappa is computed on those rows."],
    [""],
    ["Record every verdict yourself. Do not fill this sheet programmatically."],
]:
    readme.append(line)
readme.column_dimensions["A"].width = 100

out1 = OUT_DIR + r"\AUDIT_FLAGGED_307_CODING_SHEET.xlsx"
wb.save(out1)
print("written:", out1)

# ----------------------------------------------------------------- sheet 2
wb = Workbook()
ws = wb.active
ws.title = "Coding"

headers2 = [
    "ID", "Cohort", "Dataset", "Paper File", "Original Reference", "Search",
    "Verdict_Exists", "Verdict_Match", "Evidence (URL or DOI)", "Notes",
    "Matched Title", "Matched Journal", "Matched Year", "Matched DOI",
    "Double_Coded", "Coder",
]
rows2 = []
n_dbl = int(round(len(sample) * DOUBLE_CODE_SHARE))
dbl_idx = set(sample.sample(n=n_dbl, random_state=SEED).index)
for i, (idx, r) in enumerate(sample.iterrows(), start=1):
    ref = str(r["Original Reference"])
    rows2.append([
        f"V{i:03d}", r["Cohort"], r["Dataset"], r["Paper File"], ref,
        f'=HYPERLINK("{scholar_url(ref)}","search")',
        None, None, None, None,
        r["Found Title"], r["Found Journal"], r["Year"], r["DOI"],
        "YES" if idx in dbl_idx else "NO", None,
    ])

write_sheet(ws, rows2, headers2,
            input_cols=["Verdict_Exists", "Verdict_Match", "Evidence (URL or DOI)", "Notes", "Coder"],
            widths=[8, 8, 20, 30, 70, 9, 16, 20, 30, 40, 45, 25, 8, 24, 13, 10])
add_dropdown(ws, headers2, "Verdict_Exists", EXISTS_CATEGORIES, len(rows2))
add_dropdown(ws, headers2, "Verdict_Match", MATCH_CATEGORIES, len(rows2))

readme = wb.create_sheet("Protocol")
for line in [
    ["AUDIT OF THE VERIFIED STRATUM - CODING SHEET"],
    [""],
    [f"Rows: {len(rows2)}, drawn from the {len(verified):,} strings the automated pass labelled"],
    ["Verified. Proportional stratified random sample across the nine venue-year strata,"],
    [f"seed {SEED}, reproducible from code/build_audit_sheets.py."],
    ["Row order randomised, so if the audit stops early the completed rows remain an"],
    ["unbiased subsample. The rule-of-three bound applies at the sample size actually coded:"],
    ["   0 fabrications in 200 -> 95% upper bound 1.5%"],
    ["   0 fabrications in 150 -> 95% upper bound 2.0%"],
    ["   0 fabrications in 100 -> 95% upper bound 3.0%"],
    [""],
    ["WHY THIS SHEET EXISTS"],
    ["The automated pass accepted the first record Crossref returned without comparing"],
    ["title, authors or year. A fabricated reference would therefore be labelled Verified"],
    ["in almost every case. This sample is the only evidence bearing on how often that"],
    ["happened, so it carries the fabrication bound on its own."],
    [""],
    ["PROCEDURE, per row"],
    ["1. Read the string. Ignore the four Matched columns for now - they sit to the right"],
    ["   precisely so you do not read them first."],
    ["2. Search for the cited work independently. Fill Verdict_Exists:"],
    ["   Real work   - the cited work exists and the string describes it."],
    ["   Fabrication - no such work exists."],
    ["   Undecided   - state why in Notes."],
    ["3. Now read the Matched columns and fill Verdict_Match:"],
    ["   Correct match        - the retrieved record is the cited work."],
    ["   Wrong record matched - the retrieved record is a different publication."],
    ["   This column measures the precision of the automated pass and is expected to"],
    ["   produce findings independently of the fabrication question."],
    [""],
    ["DOUBLE CODING"],
    ["Rows marked YES are coded independently by a second coder. Copy the sheet, delete"],
    ["the filled columns, and give the copy to the second coder."],
    [""],
    ["Record every verdict yourself. Do not fill this sheet programmatically."],
]:
    readme.append(line)
readme.column_dimensions["A"].width = 100

out2 = OUT_DIR + r"\AUDIT_VERIFIED_SAMPLE_200_CODING_SHEET.xlsx"
wb.save(out2)
print("written:", out2)
print(f"double-coded rows: {int(round(len(rows1)*DOUBLE_CODE_SHARE))} + {int(round(len(rows2)*DOUBLE_CODE_SHARE))}")
